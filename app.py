# app.py
# -*- coding: utf-8 -*-
'''
Hebrew Appendix Maker (A4) — Streamlit
v4 updates:
- Unlimited extra appendices via “Add appendix” button (dynamic slots)
- Each extra slot can have: optional file (image/PDF), custom title, custom order
- Supports blank appendix (title only) if no file uploaded
- Keeps Hebrew intro, BiDi toggle, real pt->px sizing, DOCX RTL
'''
import io
import os
import zipfile
from dataclasses import dataclass
from typing import List, Tuple, Optional

import streamlit as st
from PIL import Image, ImageDraw, ImageFont, ImageOps, UnidentifiedImageError
import fitz  # PyMuPDF
from bidi.algorithm import get_display
from matplotlib import font_manager
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import requests

# ===== Page styling =====
st.set_page_config(page_title="יוצר נספחים A4 בעברית", page_icon="📄", layout="wide")
st.markdown(
    """
<style>
  .stApp {background: #fafafa}
  .stButton>button, .stDownloadButton>button {border-radius: 12px; font-weight: 600}
  .intro-box {background:#fff;border:1px solid #eaeaea;border-radius:14px;padding:18px 22px;margin-bottom:12px}
</style>
""",
    unsafe_allow_html=True,
)

# ===== Constants =====
DPI = 300
A4_WIDTH_INCH = 8.27
A4_HEIGHT_INCH = 11.69
A4_W = int(A4_WIDTH_INCH * DPI)
A4_H = int(A4_HEIGHT_INCH * DPI)
DEFAULT_FONT_SIZE_PT = 72
SUPPORTED_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}

HEBREW_LETTERS = ["א","ב","ג","ד","ה","ו","ז","ח","ט","י","כ","ל","מ","נ","ס","ע","פ","צ","ק","ר","ש","ת"]
BASE = len(HEBREW_LETTERS)

def index_to_hebrew(i: int) -> str:
    if i <= 0:
        return ""
    s = ""; n = i
    while n > 0:
        n -= 1
        s = HEBREW_LETTERS[n % BASE] + s
        n //= BASE
    return s

@dataclass
class FontChoice:
    name: str
    path: Optional[str]

def list_system_fonts():
    fonts = []
    for f in font_manager.fontManager.ttflist:
        try:
            fonts.append(FontChoice(name=f.name, path=f.fname))
        except Exception:
            pass
    seen, unique = set(), []
    for fc in fonts:
        if fc.name not in seen:
            seen.add(fc.name); unique.append(fc)
    unique.sort(key=lambda x: x.name.lower())
    return unique

@st.cache_data(show_spinner=False)
def fetch_david_libre() -> Optional[bytes]:
    urls = ["https://github.com/google/fonts/raw/main/ofl/davidlibre/DavidLibre-Regular.ttf"]
    for url in urls:
        try:
            r = requests.get(url, timeout=15)
            if r.status_code == 200 and r.content:
                return r.content
        except Exception:
            pass
    local_path = os.path.join("assets", "DavidLibre-Regular.ttf")
    if os.path.isfile(local_path):
        with open(local_path, "rb") as f:
            return f.read()
    return None

PT_PER_INCH = 72.0

def pt_to_px(pt: int) -> int:
    return max(1, int(round(pt * DPI / PT_PER_INCH)))

def load_truetype_font(font_bytes: bytes, size_px: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(io.BytesIO(font_bytes), size=size_px)

# ===== PDF -> images =====

def pdf_to_images(file_bytes: bytes, all_pages: bool):
    images = []
    with fitz.open(stream=file_bytes, filetype='pdf') as doc:
        pages = range(doc.page_count) if all_pages else [0]
        zoom = DPI / 72.0
        mat = fitz.Matrix(zoom, zoom)
        for pno in pages:
            page = doc.load_page(pno)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            images.append(img)
    return images


def open_uploaded_to_images(uploaded_file, all_pages: bool):
    filename = uploaded_file.name
    _, ext = os.path.splitext(filename.lower())
    data = uploaded_file.read()
    if ext == ".pdf" or uploaded_file.type == "application/pdf":
        return pdf_to_images(data, all_pages=all_pages)
    try:
        img = Image.open(io.BytesIO(data))
        img = ImageOps.exif_transpose(img).convert("RGB")
        return [img]
    except UnidentifiedImageError:
        st.error(f"קובץ {uploaded_file.name} אינו תמונה נתמכת")
        return []

# ===== Compose A4 page (content_img optional) =====

def compose_a4_page(content_img: Optional[Image.Image], title: str, font_bytes: bytes, font_size_pt: int,
                    margin_top_px: int, side_margin_px: int, spacing_px: int,
                    title_color=(0,0,0), bidi_fix=True) -> Image.Image:
    page = Image.new('RGB', (A4_W, A4_H), color=(255,255,255))
    draw = ImageDraw.Draw(page)
    title_vis = get_display(title) if bidi_fix else title
    size_px = pt_to_px(font_size_pt)
    font = load_truetype_font(font_bytes, size_px)
    tb = draw.textbbox((0,0), title_vis, font=font)
    text_w = tb[2]-tb[0]; text_h = tb[3]-tb[1]
    text_x = max(side_margin_px, (A4_W - text_w)//2)
    text_y = max(0, margin_top_px)
    draw.text((text_x, text_y), title_vis, fill=title_color, font=font)

    if content_img is None:
        # blank appendix with title only
        return page

    available_top = text_y + text_h + spacing_px
    available_height = max(1, A4_H - available_top - margin_top_px)
    available_width  = max(1, A4_W - 2*side_margin_px)
    img_w, img_h = content_img.size
    scale = min(available_width/img_w, available_height/img_h)
    new_w = max(1, int(img_w*scale)); new_h = max(1, int(img_h*scale))
    resized = content_img.resize((new_w, new_h), Image.LANCZOS)
    img_x = (A4_W - new_w)//2
    img_y = available_top + (available_height - new_h)//2
    page.paste(resized, (img_x, img_y))
    return page

# ===== DOCX =====

def add_bidi(paragraph):
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)


def create_docx(pages_info: list, font_name: str, font_size_pt: int,
                margins_mm: Tuple[float,float,float,float]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    top_mm,right_mm,bottom_mm,left_mm = margins_mm
    section.top_margin = Mm(top_mm); section.right_margin = Mm(right_mm)
    section.bottom_margin = Mm(bottom_mm); section.left_margin = Mm(left_mm)
    for pil_img, title in pages_info:
        p = doc.add_paragraph(); add_bidi(p)
        run = p.add_run(title)  # Word handles RTL
        run.font.name = font_name
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name); rFonts.set(qn('w:hAnsi'), font_name); rFonts.set(qn('w:cs'), font_name)
        run._element.rPr.append(rFonts)
        run.font.size = Pt(font_size_pt)
        bio = io.BytesIO(); pil_img.save(bio, format='PNG'); bio.seek(0)
        page_w_mm = 210 - left_mm - right_mm
        doc.add_picture(bio, width=Mm(page_w_mm))
        doc.add_page_break()
    out = io.BytesIO(); doc.save(out); return out.getvalue()

# ===== Intro =====
st.markdown(
    '<div dir="rtl" class="intro-box">\n  <h2 style="margin:0 0 8px 0">יוצר נספחים A4 בעברית</h2>\n  <p style="margin:0 0 6px 0">העלה תמונות או PDF, וקבל מסמך A4 לכל קובץ (או לכל עמוד ב-PDF), עם כותרת בעברית מעל התמונה.</p>\n  <ul style="margin:6px 0 0 0">\n    <li>שליטה בגופן: חיפוש גופנים מותקנים, העלאת TTF/OTF, או David Libre אוטומטי</li>\n    <li>גודל כותרת אמיתי בנקודות (pt), מרווחים ושוליים</li>\n    <li>יצוא: PDF מאוחד, DOCX (Word), או ZIP של נספחים בודדים</li>\n  </ul>\n</div>',
    unsafe_allow_html=True,
)

# ===== Sidebar (layout) =====
with st.sidebar:
    st.header("⚙️ הגדרות עימוד")
    sys_fonts = list_system_fonts()
    q = st.text_input("חפש גופן", value="David")
    filtered = [f for f in sys_fonts if q.lower() in f.name.lower()] if q else sys_fonts
    if not filtered:
        st.info("לא נמצא גופן תואם. אפשר להעלות TTF/OTF או להשתמש ב-David Libre.")
        filtered = sys_fonts
    selected_font_name = st.selectbox("בחר גופן מותקן", options=[f.name for f in filtered] or ["(אין)"])
    selected_font_path = next((f.path for f in sys_fonts if f.name == selected_font_name), None)
    uploaded_ttf = st.file_uploader("או העלה קובץ גופן (TTF/OTF)", type=["ttf","otf"])
    font_size_pt = st.slider("גודל כותרת (pt)", 8, 168, DEFAULT_FONT_SIZE_PT)
    margin_top = st.slider("שול עליון (px)", 0, 400, 120)
    side_margin = st.slider("שולי צד (px)", 0, 400, 80)
    spacing = st.slider("מרווח כותרת-תוכן (px)", 0, 300, 64)
    bidi_fix = st.checkbox("תקן כיוון עברית (BiDi)", value=True)
    st.subheader("מקור ה-PDF")
    pdf_mode = st.radio("עיבוד PDF", ["עמוד ראשון בלבד","כל העמודים (נספחים נפרדים)"], index=0)
    all_pages = (pdf_mode == "כל העמודים (נספחים נפרדים)")
    st.subheader("יצוא")
    export_choices = st.multiselect("פורמטי יצוא", ["PDF מאוחד","DOCX (Word)","ZIP (נספחים בודדים)"], ["PDF מאוחד"])

st.title("📄 יוצר נספחים A4 בעברית – PDF/Word")
st.caption("כותרת שחורה בתוך הדף, מעל התמונה, עם בחירת גופן וגודל אמיתי בנקודות.")
st.markdown("---")

# ===== Dynamic extra slots =====
if 'extra_slots' not in st.session_state:
    st.session_state['extra_slots'] = 0

c_add, c_reset = st.columns([0.5, 0.5])
with c_add:
    if st.button("➕ הוסף נספח נוסף"):
        st.session_state['extra_slots'] += 1
        st.experimental_rerun()
with c_reset:
    if st.button("🗑️ אפס נספחים נוספים"):
        st.session_state['extra_slots'] = 0
        st.experimental_rerun()

uploads = st.file_uploader("העלה קבצים (תמונות/‏PDF)", accept_multiple_files=True,
                           type=["pdf","png","jpg","jpeg","webp","tif","tiff"])

# ===== Collect rows (base uploads + extra slots) =====
rows = []
if uploads:
    for idx, uf in enumerate(uploads, start=1):
        default_label = f"נספח {index_to_hebrew(idx)}׳"
        c1, c2, c3 = st.columns([0.35, 0.5, 0.15])
        with c1: st.write(f"**{uf.name}**")
        with c2: title_val = st.text_input("כותרת:", value=default_label, key=f"title_{idx}")
        with c3: order_val = st.number_input("#", 1, 999999, idx, 1, key=f"order_{idx}")
        rows.append((order_val, uf, title_val))

# extra slots
base_count = len(uploads) if uploads else 0
for i in range(st.session_state['extra_slots']):
    idx = base_count + i + 1
    default_label = f"נספח {index_to_hebrew(idx)}׳"
    c1, c2, c3 = st.columns([0.35, 0.5, 0.15])
    with c1:
        extra_file = st.file_uploader("קובץ (רשות)", type=["pdf","png","jpg","jpeg","webp","tif","tiff"], key=f"extra_file_{i}")
    with c2:
        title_val = st.text_input("כותרת:", value=default_label, key=f"extra_title_{i}")
    with c3:
        order_val = st.number_input("#", 1, 999999, idx, 1, key=f"extra_order_{i}")
    rows.append((order_val, extra_file, title_val))

# ===== Process =====
rows = [r for r in rows if r[1] is not None or (r[1] is None and isinstance(r[2], str))]
rows.sort(key=lambda t: t[0])


def resolve_font_bytes():
    if uploaded_ttf is not None:
        return uploaded_ttf.read(), uploaded_ttf.name
    if selected_font_path and os.path.isfile(selected_font_path):
        with open(selected_font_path,'rb') as f:
            return f.read(), selected_font_name
    dl = fetch_david_libre()
    if dl:
        return dl, "David Libre"
    st.error("לא ניתן לטעון גופן. העלה TTF/OTF או בחר מהרשימה.")
    st.stop()

st.markdown("---")
if st.button("🔧 עבד וייצא", type="primary"):
    if not rows:
        st.warning("לא הוגדרו נספחים."); st.stop()
    font_bytes, chosen_font_name = resolve_font_bytes()
    pages = []; pages_named = []; indiv_pdf_buffers = []
    cols = st.columns(3); slot = 0

    for order_val, file_or_none, title in rows:
        images = []
        if file_or_none is not None:
            # uploaded_file from base or extra slot
            try:
                file_or_none.seek(0)
            except Exception:
                pass
            images = open_uploaded_to_images(file_or_none, all_pages=all_pages)
        else:
            images = [None]  # blank appendix

        for page_idx, img in enumerate(images, start=1):
            title_eff = title if len(images) == 1 or img is None else f"{title} – עמוד {page_idx}"
            a4 = compose_a4_page(img, title_eff, font_bytes, font_size_pt, margin_top, side_margin, spacing, bidi_fix=bidi_fix)
            pages.append(a4); pages_named.append((a4, title_eff))
            buf = io.BytesIO(); a4.save(buf, format='PDF', resolution=DPI)
            indiv_pdf_buffers.append((title_eff, buf.getvalue()))
            with cols[slot % 3]:
                preview_w = min(520, A4_W // 2); preview_h = int(preview_w * A4_H / A4_W)
                st.image(a4.resize((preview_w, preview_h)), caption=title_eff)
            slot += 1

    st.success(f"נוצרו {len(pages)} עמודים.")
    cA, cB, cC = st.columns(3)
    if "PDF מאוחד" in export_choices and pages:
        all_pdf = io.BytesIO(); pages[0].save(all_pdf, format='PDF', save_all=True, append_images=pages[1:], resolution=DPI)
        with cA: st.download_button("⬇️ הורד PDF מאוחד", all_pdf.getvalue(), "נספחים_מאוחד.pdf", "application/pdf")
    if "ZIP (נספחים בודדים)" in export_choices and indiv_pdf_buffers:
        z = io.BytesIO()
        with zipfile.ZipFile(z, 'w', compression=zipfile.ZIP_DEFLATED) as zipf:
            for title_eff, pdf_bytes in indiv_pdf_buffers:
                safe = title_eff.replace('/', '_').replace('\\', '_'); zipf.writestr(f"{safe}.pdf", pdf_bytes)
        with cB: st.download_button("⬇️ הורד ZIP של נספחים בודדים", z.getvalue(), "נספחים_בודדים.zip", "application/zip")
    if "DOCX (Word)" in export_choices and pages_named:
        docx_bytes = create_docx(pages_named, font_name=chosen_font_name, font_size_pt=font_size_pt, margins_mm=(15,15,15,15))
        with cC: st.download_button("⬇️ הורד DOCX", docx_bytes, "נספחים.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
